
from abc import ABC, abstractmethod
import os
from pdfminer.high_level import extract_text
import pdfplumber
from docx import Document
from pptx import Presentation
from pdf2image import convert_from_path
import pytesseract

# Abstract Class for File Loading
class FileLoader(ABC):
    """Abstract class defining the interface for loading various file types."""
    
    @abstractmethod
    def validate_file(self) -> bool:
        """Verify if the file is valid."""
        pass

    @abstractmethod
    def load_file(self):
        """Load the file and retrieve its content."""
        pass


# Concrete PDFLoader Class
class PDFLoader(FileLoader):
    """Implementation of FileLoader specifically for PDF files."""
    
    def __init__(self, file_path: str):
        """
        Initialize the PDFLoader with the path of a PDF file.

        Args:
            file_path (str): The full path to the PDF file.
        """
        self.file_path = file_path

    def validate_file(self) -> bool:
        """Check if the provided path points to a valid PDF file."""
        return self.file_path.endswith('.pdf') and os.path.isfile(self.file_path)

    def load_file(self):
        """
        Load the PDF file and extract its text content.

        Returns:
            str: The text extracted from the PDF file.

        Raises:
            ValueError: If the file is invalid or text extraction encounters an error.
        """
        if not self.validate_file():
            raise ValueError("Invalid PDF file")

        # Attempt text extraction using pdfminer
        try:
            text = extract_text(self.file_path)
            if text.strip():
                return text
            else:
                print("No text found in PDF using pdfminer; switching to OCR.")
        except Exception as e:
            print(f"Error during text extraction with pdfminer: {e}; switching to OCR.")

        # Use OCR as a fallback if pdfminer extraction fails
        try:
            return self.extract_text_with_ocr()
        except Exception as e:
            print(f"Error during OCR extraction: {e}")
            raise ValueError("Failed to extract text from PDF.")

    def extract_text_with_ocr(self):
        """Convert the PDF pages to images and apply OCR to extract text.

        Returns:
            str: The text obtained from the PDF images via OCR.
        """
        images = convert_from_path(self.file_path)  # Convert each page of the PDF to an image
        text = ""
        for image in images:
            try:
                ocr_text = pytesseract.image_to_string(image)  # Perform OCR to extract text
                text += ocr_text
            except Exception as e:
                print(f"Error during OCR on one of the pages: {e}")
        return text

    def extract_links(self):
        """Retrieve hyperlinks embedded in the PDF file.

        Returns:
            list: A list containing all hyperlinks found in the PDF.
        """
        links = []
        with pdfplumber.open(self.file_path) as pdf:
            for page in pdf.pages:
                if page.annots:  # Check for annotations on the page
                    for annotation in page.annots:
                        if annotation.get("uri"):
                            links.append(annotation["uri"])  # Add hyperlink to the list
        return links

    def extract_images(self):
        """Extract images embedded within the PDF file.

        Returns:
            list: A collection of image metadata from the PDF.
        """
        images = []
        with pdfplumber.open(self.file_path) as pdf:
            for page in pdf.pages:
                if page.images:
                    images.extend(page.images)  # Collect image metadata
        return images

    def extract_tables(self):
        """Extract tabular data present in the PDF.

        Returns:
            list: A list of tables obtained from the PDF.
        """
        tables = []
        with pdfplumber.open(self.file_path) as pdf:
            for page in pdf.pages:
                tables.extend(page.extract_tables())  # Gather all tables from each page
        return tables
    
    def extract_metadata(self):
        """Extract metadata information from the PDF file.

        Returns:
            dict: A dictionary containing key metadata fields of the PDF.
        """
        metadata = {}
        with pdfplumber.open(self.file_path) as pdf:
            metadata = pdf.metadata  # Fetch metadata using pdfplumber
        return metadata


# Concrete DOCXLoader Class
class DOCXLoader(FileLoader):
    """Implementation of FileLoader specifically for DOCX files."""
    
    def __init__(self, file_path: str):
        """
        Initialize the DOCXLoader with the path of a DOCX file.

        Args:
            file_path (str): The full path to the DOCX file.
        """
        self.file_path = file_path
        self.content = None

    def validate_file(self) -> bool:
        """Check if the provided path points to a valid DOCX file."""
        return self.file_path.endswith('.docx') and os.path.isfile(self.file_path)

    def load_file(self):
        """
        Load the DOCX file and retrieve its content.

        Returns:
            Document: The Document object representing the loaded DOCX file.

        Raises:
            ValueError: If the file is not a valid DOCX file.
        """
        if self.validate_file():
            self.content = Document(self.file_path)  # Load the content of the DOCX file
            return self.content
        raise ValueError("Invalid DOCX file")
    
    def extract_metadata(self):
        """Extract metadata from the DOCX file.

        Returns:
            dict: A dictionary containing the key metadata properties of the DOCX file.
        """
        metadata = {}
        doc = Document(self.file_path)
        core_properties = doc.core_properties  # Access the core properties for metadata
        metadata = {
            'title': core_properties.title,
            'author': core_properties.author,
            'subject': core_properties.subject,
            'keywords': core_properties.keywords,
            'created': core_properties.created,
            'modified': core_properties.modified,
        }
        return metadata


# Concrete PPTLoader Class
class PPTLoader(FileLoader):
    """Implementation of FileLoader specifically for PPTX files."""
    
    def __init__(self, file_path: str):
        """
        Initialize the PPTLoader with the path of a PPTX file.

        Args:
            file_path (str): The full path to the PPTX file.
        """
        self.file_path = file_path
        self.content = None

    def validate_file(self) -> bool:
        """Check if the provided path points to a valid PPTX file."""
        return self.file_path.endswith('.pptx') and os.path.isfile(self.file_path)

    def load_file(self):
        """
        Load the PPTX file and retrieve its content.

        Returns:
            Presentation: The Presentation object representing the loaded PPTX file.

        Raises:
            ValueError: If the file is not a valid PPTX file.
        """
        if self.validate_file():
            self.content = Presentation(self.file_path)  # Load the content of the PPTX file
            return self.content
        raise ValueError("Invalid PPT file")

    def extract_metadata(self):
        """Extract metadata from the PPTX file.

        Returns:
            dict: A dictionary containing the key metadata properties of the PPTX file.
        """
        metadata = {}
        presentation = Presentation(self.file_path)
        core_properties = presentation.core_properties  # Access the core properties for metadata
        metadata = {
            'title': core_properties.title,
            'author': core_properties.author,
            'subject': core_properties.subject,
            'keywords': core_properties.keywords,
            'created': core_properties.created,
            'modified': core_properties.modified,
        }
        return metadata
